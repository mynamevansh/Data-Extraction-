"""Validate the editable Page 3 Excel prototype."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src.excel_generator import generate_page_workbook
from src.layout import map_page_layout
from src.pdf_structure import extract_pdf_structure


PDF_FIXTURES = (Path("data/FR_Y9C.pdf"), Path("data/DownloadAttachment.pdf"))
OUTPUT_PATH = Path("outputs/page3_prototype.xlsx")


def main() -> None:
    pdf_path = next((path for path in PDF_FIXTURES if path.is_file()), None)
    if pdf_path is None:
        raise FileNotFoundError("Expected data/FR_Y9C.pdf or data/DownloadAttachment.pdf")

    structure = extract_pdf_structure(str(pdf_path))
    layout = map_page_layout(structure["page_structures"][2])
    result = generate_page_workbook(layout, OUTPUT_PATH)
    if not OUTPUT_PATH.is_file():
        raise AssertionError("Workbook was not generated")

    workbook = load_workbook(OUTPUT_PATH)
    if "Page 3" not in workbook.sheetnames:
        raise AssertionError("Page 3 worksheet is missing")
    worksheet = workbook["Page 3"]
    values = [cell.value for row in worksheet.iter_rows() for cell in row]
    if not any(isinstance(value, str) and "Schedule HI" in value for value in values):
        raise AssertionError("Expected native Page 3 heading in workbook")

    fields = [element for element in layout["elements"] if element["type"] == "field"]
    if len(result["field_cells"]) != len(fields) or len(fields) != 29:
        raise AssertionError("Every Page 3 form field must map to one Excel cell")
    if len(set(result["field_cells"].values())) != len(fields):
        raise AssertionError("Form fields overwrote one another")

    for name, coordinate in result["field_cells"].items():
        cell = worksheet[coordinate]
        if cell.comment is None or name not in cell.comment.text:
            raise AssertionError(f"Missing metadata comment for {name}")
        cell.value = "editable-test"
        if cell.value != "editable-test":
            raise AssertionError(f"Excel field cell is not writable: {coordinate}")

    expected_names = ("BHCK4435", "BHCK4436", "BHCKF821", "BHCK4059", "BHCK4065")
    expected_cells = [
        next(coordinate for name, coordinate in result["field_cells"].items() if field_name in name)
        for field_name in expected_names
    ]
    if len(set(expected_cells)) != len(expected_cells):
        raise AssertionError("BHCK fields must map to separate Excel cells")

    print(f"Generated: {OUTPUT_PATH.as_posix()}")
    print("Worksheet: Page 3")
    print(f"Text elements written: {result['text_elements']}")
    print(f"Form fields written: {result['form_fields']}")
    print(f"Editable cells: {result['editable_cells']}")
    print(f"Collisions resolved: {result['collisions_resolved']}")
    print("Sample mappings:")
    for name, coordinate in zip(expected_names, expected_cells):
        print(f"{name} -> {coordinate}")


if __name__ == "__main__":
    main()