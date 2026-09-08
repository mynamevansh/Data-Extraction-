"""Validate the complete 73-page PDF to Excel conversion."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from convert_full_pdf import OUTPUT_PATH, PDF_PATH, REPORT_PATH, main


def test() -> None:
    main()
    structure = json.loads(REPORT_PATH.read_text(encoding="utf-8"))
    workbook = load_workbook(OUTPUT_PATH, read_only=False)
    expected_pages = structure["pdf_pages"]
    if expected_pages != 73 or len(workbook.worksheets) != 73:
        raise AssertionError("Expected 73 PDF pages and worksheets")
    if workbook.sheetnames != [f"Page {number}" for number in range(1, 74)]:
        raise AssertionError("Worksheet names must be Page 1 through Page 73")
    if structure["pdf_form_fields"] != 1889:
        raise AssertionError("Unexpected source field count")
    if structure["pdf_form_fields"] != structure["excel_form_fields"]:
        raise AssertionError("Form fields were not fully accounted for")
    if structure["unsupported_fields"] or structure["failed_pages"]:
        raise AssertionError("Conversion report contains unsupported fields or failures")
    for worksheet in workbook.worksheets:
        if worksheet.protection.sheet or worksheet.sheet_view.showGridLines:
            raise AssertionError(f"Worksheet is not editable/clean: {worksheet.title}")
        if not any(cell.value is not None for row in worksheet.iter_rows() for cell in row):
            raise AssertionError(f"Worksheet has no mapped content: {worksheet.title}")
    page_three = workbook["Page 3"]
    if not any(cell.comment and "BHCK4435" in cell.comment.text for row in page_three.iter_rows() for cell in row):
        raise AssertionError("Page 3 representative field is missing")
    print("Full conversion test: PASS")


if __name__ == "__main__":
    test()