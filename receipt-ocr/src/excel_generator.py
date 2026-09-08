"""Generate an editable Excel prototype from a mapped PDF page."""

from __future__ import annotations

import os
import re
from collections import defaultdict
from statistics import median
from pathlib import Path
from typing import TypedDict

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter

from .layout import LayoutElement, PDFLayout


class ExcelGenerationResult(TypedDict):
    output_path: str
    text_elements: int
    form_fields: int
    editable_cells: int
    collisions_resolved: int
    field_cells: dict[str, str]


class VisualExcelGenerationResult(ExcelGenerationResult):
    visual_formatting: str


def _resolve_positions(layout: PDFLayout) -> tuple[dict[int, tuple[int, int]], int]:
    """Allocate unique cells while preserving each element's spatial order."""
    by_row: dict[int, list[tuple[int, LayoutElement]]] = defaultdict(list)
    for index, element in enumerate(layout["elements"]):
        by_row[element["excel_row"]].append((index, element))

    positions: dict[int, tuple[int, int]] = {}
    collisions = 0
    for row, row_elements in by_row.items():
        occupied: set[int] = set()
        for index, element in sorted(
            row_elements,
            key=lambda item: (item[1]["bbox"][0], item[1]["bbox"][2], item[0]),
        ):
            column = element["excel_col"]
            while column in occupied:
                column += 1
                collisions += 1
            occupied.add(column)
            positions[index] = (row, column)
    return positions, collisions


def _element_label(element: LayoutElement) -> str:
    return element["text"] if element["type"] == "text" else element["name"]


def _set_dimensions(
    worksheet, elements: list[LayoutElement], positions: dict[int, tuple[int, int]]
) -> None:
    column_metrics: dict[int, dict[str, float]] = defaultdict(
        lambda: {"source_width": 0.0, "text_width": 0.0}
    )
    row_heights: dict[int, float] = defaultdict(float)
    for index, element in enumerate(elements):
        row, column = positions[index]
        metric = column_metrics[column]
        metric["source_width"] = max(
            metric["source_width"], element["bbox"][2] - element["bbox"][0]
        )
        metric["text_width"] = max(metric["text_width"], len(_element_label(element)) * 0.8)
        line_count = max(1, element["text"].count("\n") + 1)
        row_heights[row] = max(
            row_heights[row],
            min(120.0, max(15.0, (element["bbox"][3] - element["bbox"][1]) * 1.15, line_count * 15.0)),
        )

    for column, metric in column_metrics.items():
        width = max(8.0, min(60.0, metric["source_width"] / 7.0, metric["text_width"]))
        worksheet.column_dimensions[get_column_letter(column)].width = width
    for row, height in row_heights.items():
        worksheet.row_dimensions[row].height = height


def _write_text_cell(cell, element: LayoutElement) -> None:
    cell.value = element["text"]
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.font = Font(name="Calibri", size=10)


def _write_field_cell(cell, element: LayoutElement, page: int) -> None:
    cell.value = element["value"] if element["value"] not in (None, "") else None
    cell.alignment = Alignment(horizontal="right", vertical="center")
    cell.fill = PatternFill(fill_type="solid", fgColor="EAF3F8")
    cell.border = Border(
        left=Side(style="thin", color="7F9DB9"),
        right=Side(style="thin", color="7F9DB9"),
        top=Side(style="thin", color="7F9DB9"),
        bottom=Side(style="thin", color="7F9DB9"),
    )
    cell.comment = Comment(
        "Source PDF field:\n"
        f"{element['name']}\n"
        f"PDF page: {page}\n"
        f"PDF bbox: {element['bbox']}",
        "PDF layout mapper",
    )


def generate_page_workbook(layout: PDFLayout, output_path: str | os.PathLike[str]) -> ExcelGenerationResult:
    """Generate one editable worksheet from a generic mapped PDF layout."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = f"Page {layout['page']}"
    positions, collisions = _resolve_positions(layout)

    text_count = 0
    field_count = 0
    field_cells: dict[str, str] = {}
    for index, element in enumerate(layout["elements"]):
        row, column = positions[index]
        cell = worksheet.cell(row=row, column=column)
        if element["type"] == "field":
            _write_field_cell(cell, element, layout["page"])
            field_count += 1
            field_cells[element["name"]] = cell.coordinate
        else:
            _write_text_cell(cell, element)
            text_count += 1

    _set_dimensions(worksheet, layout["elements"], positions)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return {
        "output_path": str(output),
        "text_elements": text_count,
        "form_fields": field_count,
        "editable_cells": field_count,
        "collisions_resolved": collisions,
        "field_cells": field_cells,
    }


def _visual_column_count(layout: PDFLayout) -> int:
    """Choose a compact grid from the source page width in PDF points."""
    return max(18, min(30, round(layout["page_width"] / 24)))


def _visual_column(element: LayoutElement, layout: PDFLayout, column_count: int) -> int:
    normalized_x = element["bbox"][0] / layout["page_width"]
    return max(1, min(column_count, int(normalized_x * column_count) + 1))


def _visual_row_heights(layout: PDFLayout) -> dict[int, float]:
    row_centers: dict[int, list[float]] = defaultdict(list)
    for element in layout["elements"]:
        row_centers[element["excel_row"]].append(
            (element["bbox"][1] + element["bbox"][3]) / 2
        )
    centers = {
        row: median(values) for row, values in row_centers.items()
    }
    ordered_rows = sorted(centers)
    heights: dict[int, float] = {}
    for index, row in enumerate(ordered_rows):
        if index + 1 < len(ordered_rows):
            distance = centers[ordered_rows[index + 1]] - centers[row]
        elif index:
            distance = centers[row] - centers[ordered_rows[index - 1]]
        else:
            distance = 14.0
        heights[row] = max(9.0, min(42.0, distance * 0.82))
    return heights


def _clean_visual_text(text: str) -> str:
    """Reduce repeated dotted leaders while retaining meaningful text."""
    return re.sub(r"\.{5,}", "...", text)


def _visual_font(element: LayoutElement, layout: PDFLayout) -> Font:
    height = element["bbox"][3] - element["bbox"][1]
    size = max(7.0, min(14.0, height * 0.72))
    page_fraction = element["bbox"][1] / layout["page_height"]
    bold = size >= 10.0 or page_fraction < 0.15 and size >= 8.5
    return Font(name="Arial", size=size, bold=bold, color="000000")


def _visual_field_style(cell) -> None:
    cell.fill = PatternFill(fill_type="solid", fgColor="F4F8FB")
    cell.border = Border(
        left=Side(style="thin", color="4F81BD"),
        right=Side(style="thin", color="4F81BD"),
        top=Side(style="thin", color="4F81BD"),
        bottom=Side(style="thin", color="4F81BD"),
    )
    cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=False)
    cell.protection = Protection(locked=False)


def generate_formatted_page_workbook(
    layout: PDFLayout, output_path: str | os.PathLike[str]
) -> VisualExcelGenerationResult:
    """Generate a compact, print-configured, editable workbook for one page."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = f"Page {layout['page']}"
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_LETTER
    worksheet.page_setup.orientation = worksheet.ORIENTATION_PORTRAIT
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.page_margins.top = 0.35
    worksheet.page_margins.bottom = 0.35
    worksheet.print_options.horizontalCentered = True
    column_count = _visual_column_count(layout)
    elements = list(layout["elements"])
    occupied: set[tuple[int, int]] = set()
    positions: dict[int, tuple[int, int]] = {}
    row_elements: dict[int, list[int]] = defaultdict(list)
    for index, element in enumerate(elements):
        row_elements[element["excel_row"]].append(index)

    # Reserve native field locations first. Text is moved around these cells,
    # never the other way around, so editable field mappings stay stable.
    for row, indexes in row_elements.items():
        for index in sorted(
            (index for index in indexes if elements[index]["type"] == "field"),
            key=lambda index: elements[index]["bbox"][0],
        ):
            base_column = _visual_column(elements[index], layout, column_count)
            column = base_column
            while (row, column) in occupied:
                column += 1
            occupied.add((row, column))
            positions[index] = (row, column)

    for row, indexes in row_elements.items():
        for index in sorted(
            (index for index in indexes if elements[index]["type"] != "field"),
            key=lambda index: elements[index]["bbox"][0],
        ):
            base_column = _visual_column(elements[index], layout, column_count)
            candidates = list(range(base_column, column_count + 1)) + list(
                range(base_column - 1, 0, -1)
            )
            column = next(
                (candidate for candidate in candidates if (row, candidate) not in occupied),
                column_count + 1,
            )
            while (row, column) in occupied:
                column += 1
            occupied.add((row, column))
            positions[index] = (row, column)

    max_column = max(column for _, column in positions.values())
    worksheet.print_area = f"A1:{get_column_letter(max_column)}{layout['logical_rows']}"
    column_widths: dict[int, float] = defaultdict(lambda: 2.4)
    field_cells: dict[str, str] = {}
    text_count = 0
    field_count = 0
    collisions = 0

    for index, element in enumerate(elements):
        row, column = positions[index]
        base_column = _visual_column(element, layout, column_count)
        if column != base_column:
            collisions += 1
        source_width = element["bbox"][2] - element["bbox"][0]
        text_width = len(element["text"].replace("\n", " ")) * 0.08
        column_widths[column] = max(
            column_widths[column], min(18.0, max(source_width * 0.08, text_width))
        )
        cell = worksheet.cell(row=row, column=column)

        if element["type"] == "field":
            cell.value = element["value"] if element["value"] not in (None, "") else None
            _visual_field_style(cell)
            cell.comment = Comment(
                "Source PDF field:\n"
                f"{element['name']}\n"
                f"PDF page: {layout['page']}\n"
                f"PDF bbox: {element['bbox']}",
                "PDF layout mapper",
            )
            field_cells[element["name"]] = cell.coordinate
            field_count += 1
        else:
            cell.value = _clean_visual_text(element["text"])
            cell.font = _visual_font(element, layout)
            cell.alignment = Alignment(vertical="center", wrap_text="\n" in element["text"])
            text_count += 1

    for column in range(1, column_count + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = column_widths[column]
    for row, height in _visual_row_heights(layout).items():
        worksheet.row_dimensions[row].height = height

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return {
        "output_path": str(output),
        "text_elements": text_count,
        "form_fields": field_count,
        "editable_cells": field_count,
        "collisions_resolved": collisions,
        "field_cells": field_cells,
        "visual_formatting": "PASS",
    }


def _populate_full_sheet(workbook: Workbook, layout: PDFLayout) -> dict[str, object]:
    """Populate one worksheet using the same generic visual mapping rules."""
    worksheet = workbook.create_sheet(f"Page {layout['page']}")
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_LETTER
    worksheet.page_setup.orientation = (
        worksheet.ORIENTATION_LANDSCAPE
        if layout["page_width"] > layout["page_height"]
        else worksheet.ORIENTATION_PORTRAIT
    )
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.page_margins.top = 0.35
    worksheet.page_margins.bottom = 0.35
    worksheet.print_options.horizontalCentered = True

    column_count = _visual_column_count(layout)
    elements = list(layout["elements"])
    occupied: set[tuple[int, int]] = set()
    positions: dict[int, tuple[int, int]] = {}
    row_elements: dict[int, list[int]] = defaultdict(list)
    for index, element in enumerate(elements):
        row_elements[element["excel_row"]].append(index)

    for row, indexes in row_elements.items():
        ordered = sorted(
            indexes,
            key=lambda index: (
                elements[index]["type"] != "field",
                elements[index]["bbox"][0],
            ),
        )
        for index in ordered:
            base_column = _visual_column(elements[index], layout, column_count)
            column = base_column
            while (row, column) in occupied:
                column += 1
            occupied.add((row, column))
            positions[index] = (row, column)

    max_column = max((column for _, column in positions.values()), default=1)
    worksheet.print_area = f"A1:{get_column_letter(max_column)}{max(layout['logical_rows'], 1)}"
    column_widths: dict[int, float] = defaultdict(lambda: 2.4)
    field_cells: dict[str, str] = {}
    text_count = 0
    field_count = 0
    unsupported: list[str] = []
    for index, element in enumerate(elements):
        row, column = positions[index]
        source_width = element["bbox"][2] - element["bbox"][0]
        text_width = len(element["text"].replace("\n", " ")) * 0.08
        column_widths[column] = max(
            column_widths[column], min(18.0, max(source_width * 0.08, text_width))
        )
        cell = worksheet.cell(row=row, column=column)
        if element["type"] == "field":
            field_type = element["field_type"]
            if field_type == "CheckBox" and element["value"] in (None, ""):
                cell.value = "[ ]"
            elif field_type == "Signature" and element["value"] in (None, ""):
                cell.value = "SIGNATURE"
            else:
                cell.value = element["value"] if element["value"] not in (None, "") else None
            _visual_field_style(cell)
            cell.comment = Comment(
                "Source PDF field:\n"
                f"{element['name']}\n"
                f"Field type: {field_type}\n"
                f"PDF page: {layout['page']}\n"
                f"PDF bbox: {element['bbox']}",
                "PDF layout mapper",
            )
            field_cells[element["name"]] = cell.coordinate
            field_count += 1
        else:
            cell.value = _clean_visual_text(element["text"])
            cell.font = _visual_font(element, layout)
            cell.alignment = Alignment(vertical="center", wrap_text="\n" in element["text"])
            text_count += 1

    for column in range(1, max_column + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = column_widths[column]
    for row, height in _visual_row_heights(layout).items():
        worksheet.row_dimensions[row].height = height

    return {
        "text_elements": text_count,
        "form_fields": field_count,
        "field_cells": field_cells,
        "unsupported_fields": unsupported,
        "worksheet": worksheet,
    }


def generate_full_workbook(
    structure: dict[str, object], output_path: str | os.PathLike[str]
) -> dict[str, object]:
    """Convert every extracted PDF page into one editable Excel workbook."""
    from .layout import map_page_layout

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    page_reports: dict[str, dict[str, object]] = {}
    failed_pages: list[dict[str, object]] = []
    total_text = 0
    total_fields = 0
    unsupported_fields: list[str] = []

    for page in structure["page_structures"]:
        page_number = int(page["page"])
        try:
            layout = map_page_layout(page)
            result = _populate_full_sheet(workbook, layout)
            page_reports[str(page_number)] = {
                "text_elements": result["text_elements"],
                "form_fields": result["form_fields"],
                "page_width": page["width"],
                "page_height": page["height"],
            }
            total_text += int(result["text_elements"])
            total_fields += int(result["form_fields"])
            unsupported_fields.extend(result["unsupported_fields"])
        except Exception as error:
            failed_pages.append({"page": page_number, "error": str(error)})
            page_reports[str(page_number)] = {
                "text_elements": 0,
                "form_fields": 0,
                "error": str(error),
            }

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return {
        "output_path": str(output),
        "pdf_pages": int(structure["pages"]),
        "excel_worksheets": len(workbook.worksheets),
        "text_elements": total_text,
        "excel_form_fields": total_fields,
        "unsupported_fields": unsupported_fields,
        "failed_pages": failed_pages,
        "pages": page_reports,
    }