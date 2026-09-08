"""Validate visual and editable properties of the formatted Page 3 workbook."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src.excel_generator import generate_formatted_page_workbook
from src.layout import map_page_layout
from src.pdf_structure import extract_pdf_structure


PDF_FIXTURES = (Path("data/FR_Y9C.pdf"), Path("data/DownloadAttachment.pdf"))
OUTPUT_PATH = Path("outputs/page3_formatted.xlsx")


def main() -> None:
    pdf_path = next((path for path in PDF_FIXTURES if path.is_file()), None)
    if pdf_path is None:
        raise FileNotFoundError("Expected data/FR_Y9C.pdf or data/DownloadAttachment.pdf")

    layout = map_page_layout(extract_pdf_structure(str(pdf_path))["page_structures"][2])
    result = generate_formatted_page_workbook(layout, OUTPUT_PATH)
    workbook = load_workbook(OUTPUT_PATH)
    worksheet = workbook["Page 3"]
    if worksheet.sheet_view.showGridLines:
        raise AssertionError("Gridlines must be disabled")
    if str(worksheet.page_setup.paperSize) != worksheet.PAPERSIZE_LETTER:
        raise AssertionError("Worksheet must use Letter paper")
    if worksheet.page_setup.orientation != worksheet.ORIENTATION_PORTRAIT:
        raise AssertionError("Worksheet must use portrait orientation")
    if worksheet.page_setup.fitToWidth != 1:
        raise AssertionError("Worksheet must fit to one page wide")
    if worksheet.protection.sheet:
        raise AssertionError("Worksheet must remain unprotected")

    fields = [element for element in layout["elements"] if element["type"] == "field"]
    if result["editable_cells"] != 29 or len(result["field_cells"]) != 29:
        raise AssertionError("Expected 29 editable field cells")
    for coordinate in result["field_cells"].values():
        cell = worksheet[coordinate]
        if cell.border.left.style is None or cell.fill.fill_type is None:
            raise AssertionError(f"Field cell is missing visual styling: {coordinate}")
        cell.value = "editable-test"

    values = [cell.value for row in worksheet.iter_rows() for cell in row]
    if not any(isinstance(value, str) and "Schedule HI" in value for value in values):
        raise AssertionError("Description/header text is missing")
    if not any(isinstance(value, str) and "4435" in value for value in values):
        raise AssertionError("BHCK code text is missing")
    if not any(isinstance(value, str) and "1.a.(1)(a)" in value for value in values):
        raise AssertionError("Item reference text is missing")
    if not any(
        isinstance(value, str)
        and ("provisions for credit losses" in value.lower() or "trading revenue" in value.lower())
        for value in values
    ):
        raise AssertionError("Footnote text is missing")

    expected = ("BHCK4435", "BHCK4436", "BHCKF821", "BHCK4059", "BHCK4065")
    sample_cells = [
        next(coordinate for name, coordinate in result["field_cells"].items() if field in name)
        for field in expected
    ]
    if len(set(sample_cells)) != len(sample_cells):
        raise AssertionError("BHCK fields must occupy distinct cells")

    print(f"Generated: {OUTPUT_PATH.as_posix()}")
    print("Worksheet: Page 3")
    print("Visual formatting: PASS")
    print(f"Editable fields: {result['editable_cells']}")
    print("Gridlines disabled: PASS")
    print("Letter portrait: PASS")
    print("Fit width: PASS")
    print("Sample mappings:")
    for field, coordinate in zip(expected, sample_cells):
        print(f"{field} -> {coordinate}")
    print("Visual test: PASS")


if __name__ == "__main__":
    main()